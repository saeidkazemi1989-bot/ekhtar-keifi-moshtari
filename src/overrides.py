"""اصلاح دستی گروه‌بندی از طریق فایل «اصلاحات گروه‌بندی.xlsx».

هر ردیف این فایل می‌تواند گروه/خانواده/محصول یک اخطار را بازنویسی کند.
اتصال بر اساس «سریال اخطار» است و اگر سریال خالی باشد، «شماره فنی» اعمال
می‌شود (برای اصلاح همهٔ اخطارهای یک قطعه).

ستون‌ها:
    سریال اخطار | شماره فنی | گروه محصول نهایی | خانواده محصول | محصول | توضیح

فقط ستون‌هایی که مقدار دارند جایگزین می‌شوند.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from .normalize import tech_key

COLUMNS = ["سریال اخطار", "شماره فنی", "گروه محصول نهایی", "خانواده محصول", "محصول", "توضیح"]
TARGETS = ["گروه محصول نهایی", "خانواده محصول", "محصول"]


def _text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = " ".join(str(value).split())
    return "" if text.lower() in {"nan", "none"} else text


def load_overrides(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    if not path.exists():
        return pd.DataFrame(columns=COLUMNS)
    frame = pd.read_excel(path, sheet_name=0)
    for column in COLUMNS:
        if column not in frame.columns:
            frame[column] = ""
    frame = frame[COLUMNS].copy()
    for column in COLUMNS:
        frame[column] = frame[column].map(_text)
    frame = frame[(frame["سریال اخطار"] != "") | (frame["شماره فنی"] != "")]
    return frame.reset_index(drop=True)


def apply_overrides(frame: pd.DataFrame, overrides: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """اعمال اصلاحات؛ خروجی: (داده اصلاح‌شده، تعداد ردیف‌های اصلاح‌شده)."""
    if overrides is None or overrides.empty:
        return frame, 0

    result = frame.copy()
    touched = 0
    serial_series = result["سریال اخطار"].astype(str)
    key_series = result["شماره فنی"].map(lambda v: tech_key(_text(v)))

    for _, rule in overrides.iterrows():
        serial = rule["سریال اخطار"]
        tech = rule["شماره فنی"]
        if serial:
            mask = serial_series == str(serial)
        elif tech:
            mask = key_series == tech_key(tech)
        else:
            continue
        if not mask.any():
            continue
        for target in TARGETS:
            value = rule[target]
            if value:
                result.loc[mask, target] = value
        result.loc[mask, "روش تطبیق"] = "اصلاح دستی"
        result.loc[mask, "اطمینان"] = "قطعی"
        if rule["توضیح"]:
            result.loc[mask, "یادداشت تطبیق"] = rule["توضیح"]
        result.loc[mask, "در فایل گروه‌بندی"] = [
            "بله" if str(g) not in {"", "خارج از دامنه گروه‌بندی"} else "خیر"
            for g in result.loc[mask, "گروه محصول نهایی"]
        ]
        result.loc[mask, "در دامنه گروه‌بندی"] = [
            "خیر" if str(g) == "خارج از دامنه گروه‌بندی" else "بله"
            for g in result.loc[mask, "گروه محصول نهایی"]
        ]
        touched += int(mask.sum())

    if touched:
        result["نیاز به بازبینی"] = [
            "خیر" if method == "اصلاح دستی" else flag
            for method, flag in zip(result["روش تطبیق"], result["نیاز به بازبینی"])
        ]
    return result, touched


def make_template(path: str | Path) -> Path:
    """ساخت فایل نمونهٔ اصلاحات با راهنما."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "اصلاحات"
    ws.sheet_view.rightToLeft = True
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    widths = {"A": 14, "B": 16, "C": 28, "D": 26, "E": 30, "F": 40}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    lines = [
        ("راهنمای اصلاح گروه‌بندی", ""),
        ("", ""),
        ("۱", "هر ردیف یک اصلاح است. اگر «سریال اخطار» پر شود فقط همان اخطار اصلاح می‌شود."),
        ("۲", "اگر سریال خالی و «شماره فنی» پر باشد، همهٔ اخطارهای آن شماره فنی اصلاح می‌شوند."),
        ("۳", "فقط ستون‌هایی که مقدار دارند جایگزین می‌شوند؛ بقیه دست‌نخورده می‌مانند."),
        ("۴", "مقادیر مجاز «گروه محصول نهایی»: همان ۹ گروه فایل گروه‌بندی یا «خارج از دامنه گروه‌بندی»."),
        ("۵", "پس از ذخیره، دوباره build_report.py را اجرا کنید؛ ردیف‌های اصلاح‌شده با روش «اصلاح دستی» مشخص می‌شوند."),
        ("", ""),
        ("نمونه", "سریال 211086 | خانواده: ICN1 SOREN | محصول: ICN1 SOREN | توضیح: اصلاح دستی جلوآمپر سورن"),
    ]
    for row in lines:
        guide.append(list(row))
    guide.column_dimensions["A"].width = 10
    guide.column_dimensions["B"].width = 100
    wb.save(path)
    return path
